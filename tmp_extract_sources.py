import json
import openpyxl
import re
import html
from pathlib import Path

# Extract Excel workbook structure and formulas
excel_path = Path('Coffee Grounds Data.xlsx')
wb = openpyxl.load_workbook(excel_path, data_only=False)
result = {'sheets': wb.sheetnames, 'sheet_info': {}}
for name in wb.sheetnames:
    sheet = wb[name]
    rows = []
    for i, row in enumerate(sheet.iter_rows(max_row=25, values_only=False), start=1):
        row_values = [cell.value for cell in row]
        formulas = [cell.value for cell in row if isinstance(cell.value, str) and cell.value.startswith('=')]
        rows.append({'row': i, 'values': row_values, 'formulas': formulas})
    result['sheet_info'][name] = rows
with open('excel_extract.json', 'w', encoding='utf-8') as f:
    json.dump(result, f, indent=2)

# Extract Tableau workbook calculated fields
twb_path = Path('HLF SCG Reporting.twb')
text = twb_path.read_text(encoding='utf-8')
pattern = re.compile(r"<column[^>]*caption='([^']*)'[^>]*>\s*<calculation[^>]*formula='([^']*)'", re.MULTILINE)
results = []
for m in pattern.finditer(text):
    caption = m.group(1)
    formula = html.unescape(m.group(2)).replace('&#13;&#10;', ' ').replace('&#39;', "'")
    results.append({'caption': caption, 'formula': formula})
with open('twb_extract.json', 'w', encoding='utf-8') as f:
    json.dump(results, f, indent=2)
