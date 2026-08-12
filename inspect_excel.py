import openpyxl
import os
path = 'Coffee Grounds Data - Alteryx Output.xlsx'
print('exists', os.path.exists(path))
wb = openpyxl.load_workbook(path, data_only=True)
print('sheets', wb.sheetnames)
sheet = wb[wb.sheetnames[0]]
rows = [tuple(cell.value for cell in row) for row in sheet.iter_rows(max_row=10)]
for i, row in enumerate(rows):
    print(i, row)
