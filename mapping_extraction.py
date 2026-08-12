import openpyxl
import re
import html
from pathlib import Path

source_excel = Path('Coffee Grounds Data.xlsx')
source_twb = Path('HLF SCG Reporting.twb')

wb = openpyxl.load_workbook(source_excel, data_only=False)
if 'Collection DB' not in wb.sheetnames:
    raise SystemExit('Missing Collection DB')

sheet = wb['Collection DB']
headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1, values_only=False))]
print('HEADER COUNT', len(headers))
for i, h in enumerate(headers, start=1):
    print(f'{i}: {repr(h)}')

formula_cols = set()
formula_examples = {}
for row in sheet.iter_rows(min_row=2, max_row=80, values_only=False):
    for idx, cell in enumerate(row):
        if isinstance(cell.value, str) and cell.value.startswith('='):
            formula_cols.add(idx)
            formula_examples.setdefault(idx, cell.value)

print('\nFORMULA COLUMNS', sorted(formula_cols))
for idx in sorted(formula_examples):
    print('COL', idx + 1, 'HEADER', repr(headers[idx]), 'FORMULA', formula_examples[idx])

text = source_twb.read_text(encoding='utf-8')
pattern = re.compile(r"<column[^>]*caption='([^']*)'[^>]*>\s*<calculation[^>]*formula='([^']*)'", re.MULTILINE)
seen = {}
for m in pattern.finditer(text):
    caption = html.unescape(m.group(1)).strip()
    formula = html.unescape(m.group(2)).replace('&#13;&#10;', ' ').replace('&#39;', "'").strip()
    seen[caption] = formula

print('\nTABLEAU CALC COUNT', len(seen))
for caption, formula in seen.items():
    print('---')
    print('Caption:', caption)
    print('Formula:', formula)
